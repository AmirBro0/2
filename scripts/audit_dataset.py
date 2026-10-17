"""Validate a private video dataset and its Excel labels locally."""

import argparse
from collections import defaultdict
import hashlib
import json
from pathlib import Path
import re
import stat
from zipfile import ZipFile

import cv2
from openpyxl import load_workbook


VIDEO_NAME = re.compile(r"^[XХxх]-(\d+)$")


def video_id(value: str) -> int:
    match = VIDEO_NAME.fullmatch(Path(value).stem.strip())
    if match is None:
        raise ValueError(f"Unexpected video name: {value}")
    return int(match.group(1))


def label_second(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value < 0 or int(value) != value:
            raise ValueError(f"Invalid timecode: {value}")
        return int(value)
    numbers = [int(part) for part in re.findall(r"\d+", str(value))]
    if "минут" in str(value).lower() and len(numbers) == 2:
        return numbers[0] * 60 + numbers[1]
    if len(numbers) == 1:
        return numbers[0]
    raise ValueError(f"Cannot parse timecode: {value}")


def extract_archive(archive: Path, destination: Path) -> Path:
    destination.mkdir(parents=True, exist_ok=True)
    destination = destination.resolve()
    with ZipFile(archive) as bundle:
        members = bundle.infolist()
        if len(members) > 1000 or sum(item.file_size for item in members) > 2_000_000_000:
            raise ValueError("Archive exceeds dataset limits")
        for item in members:
            target = (destination / item.filename).resolve()
            if not target.is_relative_to(destination) or stat.S_ISLNK(item.external_attr >> 16):
                raise ValueError(f"Unsafe archive entry: {item.filename}")
        damaged = bundle.testzip()
        if damaged:
            raise ValueError(f"Corrupt archive entry: {damaged}")
        bundle.extractall(destination)
    return destination


def locate_dataset(root: Path) -> tuple[Path, Path]:
    labels = list(root.rglob("open_labels.xlsx"))
    if len(labels) != 1:
        raise ValueError(f"Expected one open_labels.xlsx, found {len(labels)}")
    videos = labels[0].parent / "videos"
    if not videos.is_dir():
        raise ValueError("videos directory is missing next to labels")
    return labels[0], videos


def read_labels(path: Path) -> dict[int, dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        if len(header) < 3:
            raise ValueError("Label table needs name, timecode and verdict columns")
        labels = {}
        for row in rows:
            if not row or row[0] is None:
                continue
            number = video_id(str(row[0]))
            if number in labels:
                raise ValueError(f"Duplicate label for video {number}")
            verdict_text = str(row[2]).strip().lower()
            if verdict_text not in ("нарушение", "без нарушения"):
                raise ValueError(f"Unexpected verdict for video {number}: {row[2]}")
            timestamp = label_second(row[1])
            verdict = "violation" if verdict_text == "нарушение" else "no_violation"
            if (verdict == "violation") != (timestamp is not None):
                raise ValueError(f"Verdict/timecode mismatch for video {number}")
            labels[number] = {"verdict": verdict, "timecode_sec": timestamp}
        return labels
    finally:
        workbook.close()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def can_decode(capture: cv2.VideoCapture, second: float, fps: float, frames: int) -> bool:
    requested = min(round(second * fps), frames - 1)
    for candidate in range(requested, max(-1, requested - round(fps) - 1), -1):
        capture.set(cv2.CAP_PROP_POS_FRAMES, candidate)
        ok, frame = capture.read()
        if ok and frame is not None:
            return True
    return False


def inspect_video(path: Path, label: dict) -> dict:
    capture = cv2.VideoCapture(str(path))
    if not capture.isOpened():
        raise ValueError(f"Cannot open video: {path.name}")
    try:
        fps = capture.get(cv2.CAP_PROP_FPS)
        frames = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))
        if fps <= 0 or frames <= 0:
            raise ValueError(f"Invalid duration: {path.name}")
        duration = frames / fps
        timestamp = label["timecode_sec"]
        if timestamp is not None and timestamp > duration + 1:
            raise ValueError(f"Label is outside video: {path.name}")
        checkpoints = [0.0, duration / 2, max(0.0, duration - 0.2)]
        if timestamp is not None:
            checkpoints.append(float(timestamp))
        decode_ok = all(can_decode(capture, point, fps, frames) for point in checkpoints)
        return {
            "width": int(capture.get(cv2.CAP_PROP_FRAME_WIDTH)),
            "height": int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT)),
            "fps": round(fps, 3),
            "duration_sec": round(duration, 3),
            "bytes": path.stat().st_size,
            "sha256": file_hash(path),
            "decode_ok": decode_ok,
        }
    finally:
        capture.release()


def audit(dataset_root: Path) -> dict:
    label_path, videos_dir = locate_dataset(dataset_root)
    labels = read_labels(label_path)
    video_files = {}
    for path in videos_dir.glob("*.mov"):
        number = video_id(path.name)
        if number in video_files:
            raise ValueError(f"Duplicate video ID {number}")
        video_files[number] = path
    if set(labels) != set(video_files):
        raise ValueError(
            f"Video/label mismatch: missing video {sorted(set(labels) - set(video_files))}; "
            f"missing label {sorted(set(video_files) - set(labels))}"
        )

    videos = []
    hash_groups = defaultdict(list)
    for number, path in sorted(video_files.items()):
        item = {"id": number, "file": str(path.resolve()), **labels[number]}
        item.update(inspect_video(path, labels[number]))
        videos.append(item)
        hash_groups[item["sha256"]].append(number)
    duplicate_groups = [group for group in hash_groups.values() if len(group) > 1]
    return {
        "summary": {
            "videos": len(videos),
            "violations": sum(item["verdict"] == "violation" for item in videos),
            "no_violations": sum(item["verdict"] == "no_violation" for item in videos),
            "min_duration_sec": min(item["duration_sec"] for item in videos),
            "max_duration_sec": max(item["duration_sec"] for item in videos),
            "max_size_bytes": max(item["bytes"] for item in videos),
            "decode_failures": [item["id"] for item in videos if not item["decode_ok"]],
            "exact_duplicate_groups": duplicate_groups,
        },
        "videos": videos,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--archive", type=Path)
    source.add_argument("--dataset-dir", type=Path)
    parser.add_argument("--output", type=Path, default=Path("data/dataset_audit.json"))
    args = parser.parse_args()
    dataset_root = (
        extract_archive(args.archive, Path("data/source"))
        if args.archive
        else args.dataset_dir
    )
    result = audit(dataset_root)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
